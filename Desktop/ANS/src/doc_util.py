#!/usr/bin/env python3
"""
Document utility for converting JSON medical data to formatted documents
"""

import json
import os
from datetime import datetime
from typing import Dict, Any, List
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import OxmlElement, qn

class MedicalDocumentConverter:
    """
    Converts medical data JSON to various document formats
    """
    
    def __init__(self):
        self.timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    
    def json_to_markdown(self, json_data: Dict[str, Any], patient_name: str = "Patient") -> str:
        """
        Convert JSON medical data to Markdown format
        """
        md_content = []
        
        # Header
        md_content.append(f"# Medical Report - {patient_name}")
        md_content.append(f"**Generated on:** {self.timestamp}")
        md_content.append("")
        
        # AUTONOMICS Section
        if "AUTONOMICS" in json_data:
            md_content.append("## AUTONOMICS")
            md_content.append("")
            
            autonomics = json_data["AUTONOMICS"]
            
            # Heart Rate
            if "HR" in autonomics:
                md_content.append("### Heart Rate (HR)")
                md_content.append("| Test | Value |")
                md_content.append("|------|-------|")
                for test, value in autonomics["HR"].items():
                    md_content.append(f"| {test} | {value} |")
                md_content.append("")
            
            # Blood Pressure
            if "BP" in autonomics:
                md_content.append("### Blood Pressure (BP)")
                md_content.append("| Test | Systolic | Diastolic |")
                md_content.append("|------|----------|-----------|")
                bp_data = autonomics["BP"]
                tests = set()
                for key in bp_data.keys():
                    if "Systolic" in key:
                        tests.add(key.replace(" Systolic", ""))
                
                for test in sorted(tests):
                    systolic = bp_data.get(f"{test} Systolic", "")
                    diastolic = bp_data.get(f"{test} Diastolic", "")
                    md_content.append(f"| {test} | {systolic} | {diastolic} |")
                md_content.append("")
            
            # Other metrics
            for metric_name, metric_data in autonomics.items():
                if metric_name not in ["HR", "BP"] and isinstance(metric_data, dict):
                    md_content.append(f"### {metric_name}")
                    md_content.append("| Test | Value |")
                    md_content.append("|------|-------|")
                    for test, value in metric_data.items():
                        md_content.append(f"| {test} | {value} |")
                    md_content.append("")
        
        # CARDIOVASCULAR Section
        if "CARDIOVASCULAR" in json_data:
            md_content.append("## CARDIOVASCULAR")
            md_content.append("")
            
            cardiovascular = json_data["CARDIOVASCULAR"]
            md_content.append("| Metric | Value |")
            md_content.append("|--------|-------|")
            
            for metric, value in cardiovascular.items():
                md_content.append(f"| {metric} | {value} |")
            md_content.append("")
        
        # ELECTROCARDIOGRAM Section
        if "ELECTROCARDIOGRAM" in json_data:
            md_content.append("## ELECTROCARDIOGRAM")
            md_content.append("")
            
            ecg = json_data["ELECTROCARDIOGRAM"]
            
            for section_name, section_data in ecg.items():
                if isinstance(section_data, dict):
                    md_content.append(f"### {section_name}")
                    md_content.append("| Metric | Value |")
                    md_content.append("|--------|-------|")
                    
                    for metric, value in section_data.items():
                        md_content.append(f"| {metric} | {value} |")
                    md_content.append("")
        
        # Summary of Flagged Findings
        if "Summary_of_Flagged_Findings" in json_data:
            flagged = json_data["Summary_of_Flagged_Findings"]
            if flagged:
                md_content.append("## Summary of Flagged Findings")
                md_content.append("")
                md_content.append("⚠️ **The following values are outside normal ranges:**")
                md_content.append("")
                
                for finding in flagged:
                    md_content.append(f"- {finding}")
                md_content.append("")
        
        return "\n".join(md_content)
    
    def json_to_html(self, json_data: Dict[str, Any], patient_name: str = "Patient") -> str:
        """
        Convert JSON medical data to HTML format
        """
        html_content = []
        
        # HTML Header
        html_content.append("""
<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Medical Report - """ + patient_name + """</title>
    <style>
        body {
            font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
            line-height: 1.6;
            margin: 0;
            padding: 20px;
            background-color: #f5f5f5;
        }
        .container {
            max-width: 1200px;
            margin: 0 auto;
            background: white;
            padding: 30px;
            border-radius: 10px;
            box-shadow: 0 0 20px rgba(0,0,0,0.1);
        }
        h1 {
            color: #2c3e50;
            border-bottom: 3px solid #3498db;
            padding-bottom: 10px;
        }
        h2 {
            color: #34495e;
            margin-top: 30px;
            border-left: 4px solid #3498db;
            padding-left: 15px;
        }
        h3 {
            color: #2c3e50;
            margin-top: 25px;
        }
        table {
            width: 100%;
            border-collapse: collapse;
            margin: 15px 0;
            background: white;
        }
        th, td {
            border: 1px solid #ddd;
            padding: 12px;
            text-align: left;
        }
        th {
            background-color: #3498db;
            color: white;
            font-weight: bold;
        }
        tr:nth-child(even) {
            background-color: #f9f9f9;
        }
        .flagged {
            color: #e74c3c;
            font-weight: bold;
        }
        .summary {
            background-color: #fff3cd;
            border: 1px solid #ffeaa7;
            border-radius: 5px;
            padding: 15px;
            margin: 20px 0;
        }
        .summary h3 {
            color: #856404;
            margin-top: 0;
        }
        .timestamp {
            color: #7f8c8d;
            font-style: italic;
            margin-bottom: 20px;
        }
    </style>
</head>
<body>
    <div class="container">
""")
        
        # Header
        html_content.append(f"<h1>Medical Report - {patient_name}</h1>")
        html_content.append(f'<p class="timestamp">Generated on: {self.timestamp}</p>')
        
        # AUTONOMICS Section
        if "AUTONOMICS" in json_data:
            html_content.append('<h2>AUTONOMICS</h2>')
            
            autonomics = json_data["AUTONOMICS"]
            
            # Heart Rate
            if "HR" in autonomics:
                html_content.append('<h3>Heart Rate (HR)</h3>')
                html_content.append('<table><tr><th>Test</th><th>Value</th></tr>')
                for test, value in autonomics["HR"].items():
                    flagged_class = ' class="flagged"' if "| FLAG" in str(value) else ""
                    html_content.append(f'<tr><td>{test}</td><td{flagged_class}>{value}</td></tr>')
                html_content.append('</table>')
            
            # Blood Pressure
            if "BP" in autonomics:
                html_content.append('<h3>Blood Pressure (BP)</h3>')
                html_content.append('<table><tr><th>Test</th><th>Systolic</th><th>Diastolic</th></tr>')
                bp_data = autonomics["BP"]
                tests = set()
                for key in bp_data.keys():
                    if "Systolic" in key:
                        tests.add(key.replace(" Systolic", ""))
                
                for test in sorted(tests):
                    systolic = bp_data.get(f"{test} Systolic", "")
                    diastolic = bp_data.get(f"{test} Diastolic", "")
                    systolic_class = ' class="flagged"' if "| FLAG" in str(systolic) else ""
                    diastolic_class = ' class="flagged"' if "| FLAG" in str(diastolic) else ""
                    html_content.append(f'<tr><td>{test}</td><td{systolic_class}>{systolic}</td><td{diastolic_class}>{diastolic}</td></tr>')
                html_content.append('</table>')
            
            # Other metrics
            for metric_name, metric_data in autonomics.items():
                if metric_name not in ["HR", "BP"] and isinstance(metric_data, dict):
                    html_content.append(f'<h3>{metric_name}</h3>')
                    html_content.append('<table><tr><th>Test</th><th>Value</th></tr>')
                    for test, value in metric_data.items():
                        flagged_class = ' class="flagged"' if "| FLAG" in str(value) else ""
                        html_content.append(f'<tr><td>{test}</td><td{flagged_class}>{value}</td></tr>')
                    html_content.append('</table>')
        
        # CARDIOVASCULAR Section
        if "CARDIOVASCULAR" in json_data:
            html_content.append('<h2>CARDIOVASCULAR</h2>')
            html_content.append('<table><tr><th>Metric</th><th>Value</th></tr>')
            
            cardiovascular = json_data["CARDIOVASCULAR"]
            for metric, value in cardiovascular.items():
                flagged_class = ' class="flagged"' if "| FLAG" in str(value) else ""
                html_content.append(f'<tr><td>{metric}</td><td{flagged_class}>{value}</td></tr>')
            html_content.append('</table>')
        
        # ELECTROCARDIOGRAM Section
        if "ELECTROCARDIOGRAM" in json_data:
            html_content.append('<h2>ELECTROCARDIOGRAM</h2>')
            
            ecg = json_data["ELECTROCARDIOGRAM"]
            for section_name, section_data in ecg.items():
                if isinstance(section_data, dict):
                    html_content.append(f'<h3>{section_name}</h3>')
                    html_content.append('<table><tr><th>Metric</th><th>Value</th></tr>')
                    
                    for metric, value in section_data.items():
                        flagged_class = ' class="flagged"' if "| FLAG" in str(value) else ""
                        html_content.append(f'<tr><td>{metric}</td><td{flagged_class}>{value}</td></tr>')
                    html_content.append('</table>')
        
        # Summary of Flagged Findings
        if "Summary_of_Flagged_Findings" in json_data:
            flagged = json_data["Summary_of_Flagged_Findings"]
            if flagged:
                html_content.append('<div class="summary">')
                html_content.append('<h3>⚠️ Summary of Flagged Findings</h3>')
                html_content.append('<p>The following values are outside normal ranges:</p>')
                html_content.append('<ul>')
                
                for finding in flagged:
                    html_content.append(f'<li class="flagged">{finding}</li>')
                html_content.append('</ul>')
                html_content.append('</div>')
        
        # HTML Footer
        html_content.append("""
    </div>
</body>
</html>
""")
        
        return "\n".join(html_content)
    
    def json_to_text(self, json_data: Dict[str, Any], patient_name: str = "Patient") -> str:
        """
        Convert JSON medical data to plain text format
        """
        text_content = []
        
        # Header
        text_content.append(f"MEDICAL REPORT - {patient_name}")
        text_content.append(f"Generated on: {self.timestamp}")
        text_content.append("=" * 60)
        text_content.append("")
        
        # AUTONOMICS Section
        if "AUTONOMICS" in json_data:
            text_content.append("AUTONOMICS")
            text_content.append("-" * 20)
            text_content.append("")
            
            autonomics = json_data["AUTONOMICS"]
            
            # Heart Rate
            if "HR" in autonomics:
                text_content.append("Heart Rate (HR):")
                for test, value in autonomics["HR"].items():
                    text_content.append(f"  {test}: {value}")
                text_content.append("")
            
            # Blood Pressure
            if "BP" in autonomics:
                text_content.append("Blood Pressure (BP):")
                bp_data = autonomics["BP"]
                tests = set()
                for key in bp_data.keys():
                    if "Systolic" in key:
                        tests.add(key.replace(" Systolic", ""))
                
                for test in sorted(tests):
                    systolic = bp_data.get(f"{test} Systolic", "")
                    diastolic = bp_data.get(f"{test} Diastolic", "")
                    text_content.append(f"  {test}: {systolic}/{diastolic}")
                text_content.append("")
            
            # Other metrics
            for metric_name, metric_data in autonomics.items():
                if metric_name not in ["HR", "BP"] and isinstance(metric_data, dict):
                    text_content.append(f"{metric_name}:")
                    for test, value in metric_data.items():
                        text_content.append(f"  {test}: {value}")
                    text_content.append("")
        
        # CARDIOVASCULAR Section
        if "CARDIOVASCULAR" in json_data:
            text_content.append("CARDIOVASCULAR")
            text_content.append("-" * 20)
            text_content.append("")
            
            cardiovascular = json_data["CARDIOVASCULAR"]
            for metric, value in cardiovascular.items():
                text_content.append(f"{metric}: {value}")
            text_content.append("")
        
        # ELECTROCARDIOGRAM Section
        if "ELECTROCARDIOGRAM" in json_data:
            text_content.append("ELECTROCARDIOGRAM")
            text_content.append("-" * 20)
            text_content.append("")
            
            ecg = json_data["ELECTROCARDIOGRAM"]
            for section_name, section_data in ecg.items():
                if isinstance(section_data, dict):
                    text_content.append(f"{section_name}:")
                    for metric, value in section_data.items():
                        text_content.append(f"  {metric}: {value}")
                    text_content.append("")
        
        # Summary of Flagged Findings
        if "Summary_of_Flagged_Findings" in json_data:
            flagged = json_data["Summary_of_Flagged_Findings"]
            if flagged:
                text_content.append("SUMMARY OF FLAGGED FINDINGS")
                text_content.append("-" * 30)
                text_content.append("")
                text_content.append("⚠️ The following values are outside normal ranges:")
                text_content.append("")
                
                for finding in flagged:
                    text_content.append(f"• {finding}")
                text_content.append("")
        
        return "\n".join(text_content)
    
    def json_to_docx(self, json_data: Dict[str, Any], patient_name: str = "Patient") -> Document:
        """
        Convert JSON medical data to DOCX format with black text and new page headings
        """
        doc = Document()
        
        # Set document margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
        
        def add_black_text(paragraph, text, size=Pt(10), bold=False):
            """Add black text to paragraph"""
            run = paragraph.add_run(text)
            run.font.color.rgb = None  # Default black
            run.font.size = size
            run.font.name = 'Arial'
            run.bold = bold
            return run
        
        def add_text_with_flag_formatting(paragraph, text, size=Pt(10)):
            """Add text with proper FLAG formatting (bold FLAG)"""
            if "| FLAG" in text:
                # Split text and FLAG
                parts = text.split("| FLAG")
                main_text = parts[0].strip()
                flag_text = "| FLAG"  # No extra spaces before FLAG
                
                # Add main text
                main_run = paragraph.add_run(main_text)
                main_run.font.color.rgb = None  # Black
                main_run.font.size = size
                main_run.font.name = 'Arial'
                main_run.bold = False
                
                # Add FLAG text in bold
                flag_run = paragraph.add_run(flag_text)
                flag_run.font.color.rgb = None  # Black
                flag_run.font.size = size
                flag_run.font.name = 'Arial'
                flag_run.bold = True
            else:
                # No FLAG, just add normal text
                run = paragraph.add_run(text)
                run.font.color.rgb = None  # Black
                run.font.size = size
                run.font.name = 'Arial'
                run.bold = False
        
        def add_main_heading(doc, text, is_first=False):
            """Add main heading with new page (except first one)"""
            if not is_first:
                # Add page break before main heading (except first)
                doc.add_page_break()
            
            # Create paragraph instead of heading to avoid default styling
            para = doc.add_paragraph()
            para.style = 'Normal'
            
            # Add text with explicit black formatting
            run = para.add_run(text)
            run.font.size = Pt(14)
            run.font.name = 'Arial'
            run.font.color.rgb = None  # Black
            run.bold = True
            
            return para
        
        # AUTONOMICS Section (first section - no page break)
        if "AUTONOMICS" in json_data:
            add_main_heading(doc, "AUTONOMICS //", is_first=True)
            
            autonomics = json_data["AUTONOMICS"]
            
            # Heart Rate
            if "HR" in autonomics:
                hr_para = doc.add_paragraph()
                hr_para.style = 'Normal'
                hr_run = hr_para.add_run('HR')
                hr_run.font.size = Pt(12)
                hr_run.font.name = 'Arial'
                hr_run.font.color.rgb = None  # Black
                hr_run.bold = True
                
                for test, value in autonomics["HR"].items():
                    para = doc.add_paragraph()
                    add_text_with_flag_formatting(para, f'• {test}: {value}')
            
            # Blood Pressure
            if "BP" in autonomics:
                bp_para = doc.add_paragraph()
                bp_para.style = 'Normal'
                bp_run = bp_para.add_run('BP')
                bp_run.font.size = Pt(12)
                bp_run.font.name = 'Arial'
                bp_run.font.color.rgb = None  # Black
                bp_run.bold = True
                
                bp_data = autonomics["BP"]
                
                # Check if BP data is in separated format (Systolic/Diastolic) or combined format (126/78)
                has_separated_format = any("Systolic" in key for key in bp_data.keys())
                
                if has_separated_format:
                    # Handle separated format (Systolic/Diastolic keys)
                    tests = set()
                    for key in bp_data.keys():
                        if "Systolic" in key:
                            tests.add(key.replace(" Systolic", ""))
                    
                    for test in sorted(tests):
                        systolic = bp_data.get(f"{test} Systolic", "")
                        diastolic = bp_data.get(f"{test} Diastolic", "")
                        para = doc.add_paragraph()
                        add_text_with_flag_formatting(para, f'• {test} Systolic: {systolic}')
                        para = doc.add_paragraph()
                        add_text_with_flag_formatting(para, f'• {test} Diastolic: {diastolic}')
                else:
                    # Handle combined format (126/78) - convert to separated format for AUTONOMICS
                    for test, bp_value in bp_data.items():
                        # Parse BP value to separate systolic and diastolic
                        if '/' in str(bp_value):
                            try:
                                systolic, diastolic = str(bp_value).split('/', 1)
                                systolic = systolic.strip()
                                diastolic = diastolic.strip()
                                para = doc.add_paragraph()
                                add_text_with_flag_formatting(para, f'• {test} Systolic: {systolic}')
                                para = doc.add_paragraph()
                                add_text_with_flag_formatting(para, f'• {test} Diastolic: {diastolic}')
                            except:
                                # Fallback to original format if parsing fails
                                para = doc.add_paragraph()
                                add_text_with_flag_formatting(para, f'• {test}: {bp_value}')
                        else:
                            # Fallback to original format if no separator found
                            para = doc.add_paragraph()
                            add_text_with_flag_formatting(para, f'• {test}: {bp_value}')
            
            # Other metrics
            for metric_name, metric_data in autonomics.items():
                if metric_name not in ["HR", "BP"] and isinstance(metric_data, dict):
                    metric_para = doc.add_paragraph()
                    metric_para.style = 'Normal'
                    metric_run = metric_para.add_run(metric_name)
                    metric_run.font.size = Pt(12)
                    metric_run.font.name = 'Arial'
                    metric_run.font.color.rgb = None  # Black
                    metric_run.bold = True
                    
                    for test, value in metric_data.items():
                        para = doc.add_paragraph()
                        add_text_with_flag_formatting(para, f'• {test}: {value}')
        
        # CARDIOVASCULAR Section
        if "CARDIOVASCULAR" in json_data:
            add_main_heading(doc, "CARDIOVASCULAR //")
            
            cardiovascular = json_data["CARDIOVASCULAR"]
            for metric, value in cardiovascular.items():
                para = doc.add_paragraph()
                add_text_with_flag_formatting(para, f'• {metric}: {value}')
        
        # BLOOD_PRESSURE_SUMMARY Section
        if "BLOOD_PRESSURE_SUMMARY" in json_data:
            add_main_heading(doc, "BLOOD PRESSURE SUMMARY //")
            
            bp_summary = json_data["BLOOD_PRESSURE_SUMMARY"]
            
            # Group by test name
            test_groups = {}
            for key, value in bp_summary.items():
                if "Systolic" in key:
                    test_name = key.replace(" Systolic", "")
                    if test_name not in test_groups:
                        test_groups[test_name] = {}
                    test_groups[test_name]["systolic"] = value
                elif "Diastolic" in key:
                    test_name = key.replace(" Diastolic", "")
                    if test_name not in test_groups:
                        test_groups[test_name] = {}
                    test_groups[test_name]["diastolic"] = value
                else:
                    # Fallback for unparsed values
                    test_groups[key] = {"raw": value}
            
            # Display in bullet format
            for test_name, values in test_groups.items():
                if "raw" in values:
                    para = doc.add_paragraph()
                    add_text_with_flag_formatting(para, f'• {test_name}: {values["raw"]}')
                else:
                    systolic = values.get("systolic", "")
                    diastolic = values.get("diastolic", "")
                    if systolic and diastolic:
                        para = doc.add_paragraph()
                        add_text_with_flag_formatting(para, f'• {test_name} Systolic: {systolic}')
                        para = doc.add_paragraph()
                        add_text_with_flag_formatting(para, f'• {test_name} Diastolic: {diastolic}')

        # ELECTROCARDIOGRAM Section
        if "ELECTROCARDIOGRAM" in json_data:
            add_main_heading(doc, "ELECTROCARDIOGRAM //")
            
            ecg = json_data["ELECTROCARDIOGRAM"]
            for section_name, section_data in ecg.items():
                if isinstance(section_data, dict):
                    section_para = doc.add_paragraph()
                    section_para.style = 'Normal'
                    section_run = section_para.add_run(section_name)
                    section_run.font.size = Pt(12)
                    section_run.font.name = 'Arial'
                    section_run.font.color.rgb = None  # Black
                    section_run.bold = True
                    
                    # Add data with bullet points
                    for metric, value in section_data.items():
                        para = doc.add_paragraph()
                        add_text_with_flag_formatting(para, f'• {metric}: {value}')
        
        # Summary of Flagged Findings
        if "Summary_of_Flagged_Findings" in json_data:
            add_main_heading(doc, "Summary of Flagged Findings")
            
            flagged = json_data["Summary_of_Flagged_Findings"]
            if flagged:
                for finding in flagged:
                    para = doc.add_paragraph()
                    add_text_with_flag_formatting(para, f'• {finding}')
        
        return doc
    
    def convert_json_file(self, json_file_path: str, output_format: str = "markdown", 
                        patient_name: str = "Patient", output_dir: str = "documents") -> str:
        """
        Convert JSON file to document format
        
        Args:
            json_file_path: Path to the JSON file
            output_format: Format to convert to (markdown, html, text)
            patient_name: Name of the patient
            output_dir: Directory to save the output file
        
        Returns:
            Path to the generated document file
        """
        # Ensure output directory exists
        os.makedirs(output_dir, exist_ok=True)
        
        # Load JSON data
        with open(json_file_path, 'r', encoding='utf-8') as f:
            json_data = json.load(f)
        
        # Generate filename - fixed name for single output
        filename = "medical_report"
        
        # Convert based on format
        if output_format.lower() == "markdown":
            content = self.json_to_markdown(json_data, patient_name)
            file_path = os.path.join(output_dir, f"{filename}.md")
            file_extension = "md"
            # Write the file
            with open(file_path, 'w', encoding='utf-8') as f:
                f.write(content)
        elif output_format.lower() == "html":
            content = self.json_to_html(json_data, patient_name)
            file_path = os.path.join(output_dir, f"{filename}.html")
            file_extension = "html"
            # Write the file
            with open(file_path, 'w', encoding='utf-8') as f:
                f.write(content)
        elif output_format.lower() == "text":
            content = self.json_to_text(json_data, patient_name)
            file_path = os.path.join(output_dir, f"{filename}.txt")
            file_extension = "txt"
            # Write the file
            with open(file_path, 'w', encoding='utf-8') as f:
                f.write(content)
        elif output_format.lower() == "docx":
            doc = self.json_to_docx(json_data, patient_name)
            file_path = os.path.join(output_dir, f"{filename}.docx")
            file_extension = "docx"
            # Save the document
            doc.save(file_path)
        else:
            raise ValueError(f"Unsupported format: {output_format}. Supported formats: markdown, html, text, docx")
        
        print(f"✅ Document generated: {file_path}")
        return file_path

def main():
    """
    Main function to demonstrate usage
    """
    converter = MedicalDocumentConverter()
    
    # Example usage
    json_file = "json_out/enhanced_extraction.json"
    
    if os.path.exists(json_file):
        print("🔄 Converting JSON to documents...")
        
        # Convert to DOCX only
        try:
            docx_path = converter.convert_json_file(
                json_file_path=json_file,
                output_format="docx",
                patient_name="John Doe",
                output_dir="documents"
            )
            print(f"📄 DOCX document: {docx_path}")
        except Exception as e:
            print(f"❌ Error converting to DOCX: {e}")
    else:
        print(f"❌ JSON file not found: {json_file}")
        print("💡 Please run the PDF processing pipeline first to generate the JSON file")

if __name__ == "__main__":
    main()
